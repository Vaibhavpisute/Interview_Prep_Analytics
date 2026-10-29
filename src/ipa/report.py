"""Weekly Excel report (openpyxl) with conditional formatting."""

from __future__ import annotations

import logging
from datetime import UTC, datetime
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, DataBarRule, FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy.engine import Engine

log = logging.getLogger(__name__)

SHEETS = ["Summary", "Trends", "Weak Areas", "Anomalies"]
HEADER_FILL = PatternFill("solid", start_color="1F3A5F")
HEADER_FONT = Font(bold=True, color="FFFFFF")
RED = PatternFill("solid", start_color="F8CBAD", end_color="F8CBAD")
AMBER = PatternFill("solid", start_color="FFE699", end_color="FFE699")
YELLOW = PatternFill("solid", start_color="FFF2CC", end_color="FFF2CC")
GREEN = PatternFill("solid", start_color="C6E0B4", end_color="C6E0B4")

ANOMALY_SQL = """
SELECT a.anomaly_id, a.rule, a.severity, s.session_key, a.category, a.answer_id,
       a.metric_value, a.detail, a.detected_at
FROM anomalies a
LEFT JOIN dim_session s ON s.session_id = a.session_id
ORDER BY CASE a.severity WHEN 'high' THEN 0 WHEN 'medium' THEN 1 ELSE 2 END, a.anomaly_id
"""


def _q(engine: Engine, sql: str) -> pd.DataFrame:
    return pd.read_sql(sql, engine)


def _write_table(ws, df: pd.DataFrame, start_row: int = 1) -> None:
    for j, col in enumerate(df.columns, start=1):
        c = ws.cell(row=start_row, column=j, value=str(col))
        c.fill, c.font = HEADER_FILL, HEADER_FONT
        c.alignment = Alignment(horizontal="center")
    for i, row in enumerate(df.itertuples(index=False), start=start_row + 1):
        for j, val in enumerate(row, start=1):
            if pd.isna(val):
                val = None
            elif isinstance(val, pd.Timestamp):
                val = val.to_pydatetime()
            elif hasattr(val, "item"):
                val = val.item()
            ws.cell(row=i, column=j, value=val)
    for j, col in enumerate(df.columns, start=1):
        longest = max([len(str(col))] + [len(str(v)) for v in df.iloc[:, j - 1].head(200)])
        ws.column_dimensions[get_column_letter(j)].width = min(max(12, longest + 2), 80)
    ws.freeze_panes = ws.cell(row=start_row + 1, column=1)


def collect(engine: Engine) -> dict[str, pd.DataFrame]:
    """Read everything the report needs (also used by the dashboard and tests)."""
    trend = _q(engine, "SELECT * FROM v_score_trend_by_category ORDER BY week_start, category")
    pivot = trend.pivot(index="week_start", columns="category", values="avg_score").round(2)
    weekly = _q(engine, "SELECT * FROM v_weekly_delivery ORDER BY week_start").set_index(
        "week_start"
    )
    overall = _q(
        engine,
        "SELECT d.week_start, AVG(f.score * 1.0) AS overall_score FROM fact_answers f "
        "JOIN dim_date d ON d.date_id = f.date_id WHERE f.score IS NOT NULL GROUP BY d.week_start",
    ).set_index("week_start")
    trends = pivot.join(overall).join(weekly[["avg_wpm", "fillers_per_100_words", "answers"]])
    trends = trends.round(2).reset_index()
    trends["week_start"] = trends["week_start"].astype(str)
    weak = _q(engine, "SELECT * FROM v_weakest_categories")
    weak["avg_score"] = weak["avg_score"].round(2)
    weak.insert(0, "rank", range(1, len(weak) + 1))
    return {"trends": trends, "weak": weak, "anomalies": _q(engine, ANOMALY_SQL)}


def build_report(engine: Engine, out_path: Path) -> Path:
    data = collect(engine)
    trends, weak, anom = data["trends"], data["weak"], data["anomalies"]
    counts = _q(
        engine,
        "SELECT (SELECT COUNT(*) FROM dim_session) AS sessions, "
        "(SELECT COUNT(*) FROM fact_answers) AS answers, "
        "(SELECT COUNT(*) FROM fact_answers WHERE score IS NOT NULL) AS scored, "
        "(SELECT AVG(score * 1.0) FROM fact_answers) AS avg_score, "
        "(SELECT AVG(wpm * 1.0) FROM fact_answers) AS avg_wpm, "
        "(SELECT COALESCE(SUM(rows_rejected), 0) FROM load_log) AS rejected",
    ).iloc[0]
    repeat = _q(engine, "SELECT * FROM v_repeat_rate_summary").iloc[0]
    filler_rate = _q(
        engine,
        "SELECT 100.0 * SUM(filler_count) / NULLIF(SUM(word_count), 0) AS r FROM fact_answers",
    ).iloc[0, 0]

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = "Interview Prep Analytics - weekly report"
    ws["A1"].font = Font(bold=True, size=16)
    ws["A2"] = f"Generated {datetime.now(UTC):%Y-%m-%d %H:%M} UTC. Practice data only."
    latest = trends.dropna(subset=["overall_score"]).tail(1)
    prior = trends.dropna(subset=["overall_score"]).tail(2).head(1)
    latest_score = float(latest["overall_score"].iloc[0]) if len(latest) else None
    change = (
        latest_score - float(prior["overall_score"].iloc[0])
        if latest_score is not None and len(prior)
        else None
    )
    metrics = [
        ("Practice sessions", int(counts["sessions"])),
        ("Answers loaded", int(counts["answers"])),
        ("Answers scored", int(counts["scored"])),
        ("Average score (1-10)", _r(counts["avg_score"])),
        ("Latest week average score", _r(latest_score)),
        ("Change vs previous week", _r(change)),
        ("Weakest category", weak["category"].iloc[0] if len(weak) else None),
        ("Average pace (words per minute)", _r(counts["avg_wpm"], 0)),
        ("Filler words per 100 words", _r(filler_rate)),
        ("Question repeat rate (%)", _r(repeat["repeat_rate_pct"], 1)),
        ("Rows rejected at load", int(counts["rejected"])),
        ("Anomalies found", len(anom)),
    ]
    for i, (label, value) in enumerate(metrics, start=4):
        ws.cell(row=i, column=1, value=label).font = Font(bold=True)
        ws.cell(row=i, column=2, value=value)
    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 22
    anom_row = 4 + len(metrics) - 1
    ws.conditional_formatting.add(
        f"B{anom_row}", CellIsRule(operator="greaterThan", formula=["0"], fill=AMBER)
    )
    ws.conditional_formatting.add("B9", CellIsRule(operator="lessThan", formula=["0"], fill=RED))
    ws.conditional_formatting.add(
        "B9", CellIsRule(operator="greaterThan", formula=["0"], fill=GREEN)
    )

    wt = wb.create_sheet("Trends")
    _write_table(wt, trends)
    if len(trends):
        last = len(trends) + 1
        n_cat = max(len(trends.columns) - 5, 0)
        if n_cat:
            rng = f"B2:{get_column_letter(1 + n_cat + 1)}{last}"  # categories + overall score
            wt.conditional_formatting.add(
                rng,
                ColorScaleRule(
                    start_type="num", start_value=1, start_color="F8696B",
                    mid_type="num", mid_value=5.5, mid_color="FFEB84",
                    end_type="num", end_value=10, end_color="63BE7B",
                ),
            )  # fmt: skip
        wpm_col = get_column_letter(list(trends.columns).index("avg_wpm") + 1)
        wt.conditional_formatting.add(
            f"{wpm_col}2:{wpm_col}{last}",
            FormulaRule(formula=[f"OR({wpm_col}2<90,{wpm_col}2>190)"], fill=RED),
        )

    ww = wb.create_sheet("Weak Areas")
    _write_table(ww, weak)
    if len(weak):
        last = len(weak) + 1
        ww.conditional_formatting.add(
            f"C2:C{last}", CellIsRule(operator="lessThan", formula=["6"], fill=RED)
        )
        ww.conditional_formatting.add(
            f"C2:C{last}",
            DataBarRule(start_type="num", start_value=0, end_type="num", end_value=10,
                        color="5B9BD5"),
        )  # fmt: skip

    wa = wb.create_sheet("Anomalies")
    _write_table(wa, anom)
    if len(anom):
        last = len(anom) + 1
        for sev, fill in (("high", RED), ("medium", AMBER), ("low", YELLOW)):
            wa.conditional_formatting.add(
                f"A2:I{last}", FormulaRule(formula=[f'$C2="{sev}"'], fill=fill)
            )

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    log.info("Wrote %s (%d trend weeks, %d anomalies)", out_path, len(trends), len(anom))
    return out_path


def _r(value: object, digits: int = 2) -> float | int | None:
    if value is None or pd.isna(value):
        return None
    return round(float(value), digits) if digits else int(round(float(value)))
