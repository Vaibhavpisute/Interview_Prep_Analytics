from __future__ import annotations

import json

import pytest
from openpyxl import load_workbook
from typer.testing import CliRunner

from conftest import SAMPLE_RAW, make_row
from ipa.cli import app

runner = CliRunner()


@pytest.fixture
def env(monkeypatch, tmp_path):
    db = tmp_path / "cli.db"
    monkeypatch.setenv("DATABASE_URL", f"sqlite:///{db.as_posix()}")
    monkeypatch.setenv("LLM_PROVIDER", "fake")
    for key in ("ANTHROPIC_API_KEY", "OPENAI_API_KEY", "GEMINI_API_KEY"):
        monkeypatch.setenv(key, "")
    monkeypatch.chdir(tmp_path)
    return tmp_path


def invoke(*args):
    result = runner.invoke(app, list(args))
    assert result.exit_code == 0, result.output
    return result.output


def test_help_lists_every_command():
    out = invoke("--help")
    for cmd in ("init-db", "ingest", "generate", "score", "check", "weekly-report", "demo"):
        assert cmd in out
    assert "export" in out


def test_step_by_step_pipeline(env):
    invoke("init-db")
    out = invoke("ingest", "--raw-dir", str(SAMPLE_RAW), "--rejected-dir", str(env / "rej"))
    assert "sessions_2026_08.csv" in out and "loaded" in out
    again = invoke("ingest", "--raw-dir", str(SAMPLE_RAW), "--rejected-dir", str(env / "rej"))
    assert again.count("skipped") >= 3
    assert "Scored" in invoke("score")
    assert "total" in invoke("check")
    invoke("weekly-report", "--out", str(env / "r.xlsx"))
    assert load_workbook(env / "r.xlsx").sheetnames == [
        "Summary", "Trends", "Weak Areas", "Anomalies",
    ]  # fmt: skip
    invoke("export", "--out-dir", str(env / "csv"))
    assert (env / "csv" / "fact_answers.csv").exists()


def test_generate_writes_a_practice_sheet(env):
    jd = env / "jd.txt"
    jd.write_text("Job title: Data Engineer\nSQL and Spark.", encoding="utf-8")
    out = invoke("generate", "--jd", str(jd), "--n", "4", "--out-dir", str(env / "gen"))
    assert "4 practice questions" in out
    sheet = next((env / "gen").glob("practice_*.json"))
    assert len(json.loads(sheet.read_text(encoding="utf-8"))) == 4


def test_generate_with_missing_jd_gives_a_clear_error(env):
    result = runner.invoke(app, ["generate", "--jd", "nope.txt"])
    assert result.exit_code == 1 and "not found" in result.output


def test_score_without_api_key_explains_what_to_do(env):
    result = runner.invoke(app, ["score", "--provider", "anthropic"])
    assert result.exit_code == 1
    assert "ANTHROPIC_API_KEY" in result.output and "--provider fake" in result.output


def test_ingest_cli_with_custom_folder(env):
    raw = env / "raw"
    raw.mkdir()
    (raw / "a.json").write_text(json.dumps([make_row(1)]), encoding="utf-8")
    assert "1" in invoke("ingest", "--raw-dir", str(raw), "--rejected-dir", str(env / "rej"))


def test_demo_runs_end_to_end_with_the_fake_llm(env):
    out = invoke(
        "demo", "--reset", "--db", f"sqlite:///{(env / 'demo.db').as_posix()}",
        "--raw-dir", str(SAMPLE_RAW), "--out-dir", str(env / "out"),
        "--jd", str(SAMPLE_RAW.parent / "jd" / "sample_jd.txt"),
    )  # fmt: skip
    assert "96 answers loaded" in out and "16 anomalies" in out
    assert (env / "out" / "weekly_report.xlsx").exists()
    assert (env / "out" / "export" / "fact_answers.csv").exists()
