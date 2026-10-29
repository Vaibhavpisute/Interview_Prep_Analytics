from __future__ import annotations

import pandas as pd
import pytest
from openpyxl import load_workbook

from ipa.checks import run_checks
from ipa.export import EXPORT_TABLES, export_csv
from ipa.llm.client import LLMClient
from ipa.llm.fake import FakeProvider
from ipa.report import SHEETS, build_report, collect
from ipa.workflow import score_pending


@pytest.fixture
def ready(loaded_engine):
    score_pending(loaded_engine, LLMClient(FakeProvider(), sleep=lambda s: None, max_retries=3))
    run_checks(loaded_engine)
    return loaded_engine


def test_report_has_the_four_sheets_with_expected_row_counts(ready, tmp_path):
    path = build_report(ready, tmp_path / "out" / "weekly.xlsx")
    wb = load_workbook(path)
    assert wb.sheetnames == SHEETS == ["Summary", "Trends", "Weak Areas", "Anomalies"]
    data = collect(ready)
    n_weeks = pd.read_sql("SELECT COUNT(DISTINCT week_start) AS n FROM v_weekly_delivery", ready)
    assert wb["Trends"].max_row == 1 + int(n_weeks["n"].iloc[0]) == 11
    assert wb["Weak Areas"].max_row == 1 + 6
    assert wb["Anomalies"].max_row == 1 + len(data["anomalies"]) == 17


def test_summary_sheet_numbers_match_the_database(ready, tmp_path):
    ws = load_workbook(build_report(ready, tmp_path / "w.xlsx"))["Summary"]
    values = {ws.cell(row=r, column=1).value: ws.cell(row=r, column=2).value for r in range(4, 20)}
    assert values["Practice sessions"] == 30
    assert values["Answers loaded"] == 96
    assert values["Answers scored"] == 94
    assert values["Rows rejected at load"] == 4
    assert values["Anomalies found"] == 16
    assert 1 < values["Average score (1-10)"] < 10


def test_conditional_formatting_is_applied(ready, tmp_path):
    wb = load_workbook(build_report(ready, tmp_path / "w.xlsx"))
    for name in SHEETS:
        assert len(wb[name].conditional_formatting) > 0, name
    kinds = {r.type for cf in wb["Weak Areas"].conditional_formatting for r in cf.rules}
    assert {"cellIs", "dataBar"} <= kinds
    assert "colorScale" in {r.type for cf in wb["Trends"].conditional_formatting for r in cf.rules}


def test_trends_sheet_columns_and_header_style(ready, tmp_path):
    ws = load_workbook(build_report(ready, tmp_path / "w.xlsx"))["Trends"]
    header = [c.value for c in ws[1]]
    assert header[0] == "week_start"
    assert {"SQL", "Python", "Spark", "ETL", "Data Modelling", "Behavioural"} <= set(header)
    assert {"overall_score", "avg_wpm", "fillers_per_100_words"} <= set(header)
    assert ws["A1"].font.bold and ws.freeze_panes == "A2"


def test_report_on_an_empty_database_still_has_all_sheets(engine, tmp_path):
    wb = load_workbook(build_report(engine, tmp_path / "empty.xlsx"))
    assert wb.sheetnames == SHEETS
    assert wb["Anomalies"].max_row == 1


def test_export_writes_one_csv_per_table(ready, tmp_path):
    paths = export_csv(ready, tmp_path / "csv")
    assert [p.stem for p in paths] == EXPORT_TABLES
    fact = pd.read_csv(tmp_path / "csv" / "fact_answers.csv")
    assert len(fact) == 96 and {"answer_id", "score", "wpm"} <= set(fact.columns)
    assert len(pd.read_csv(tmp_path / "csv" / "anomalies.csv")) == 16
